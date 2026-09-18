import io
import pytest
from httpx import AsyncClient, ASGITransport
from docx import Document
from reportlab.pdfgen import canvas

from src.api import app
from src.agent_1_resume_parser import (
    extract_text_from_pdf_bytes,
    extract_text_from_docx_bytes,
    extract_text_from_txt_bytes,
    parse_resume_from_bytes,
    CandidateClaims,
)


def create_sample_pdf() -> bytes:
    buf = io.BytesIO()
    c = canvas.Canvas(buf)
    c.drawString(100, 750, "Jane Doe")
    c.drawString(100, 735, "Email: jane.doe@example.com | Phone: +1-555-0199 | Location: San Francisco, CA")
    c.drawString(100, 720, "Role: Senior Backend Engineer")
    c.drawString(100, 700, "Professional Summary:")
    c.drawString(100, 685, "Experienced backend engineer with 7+ years developing distributed microservices.")
    c.drawString(100, 665, "Technical Skills:")
    c.drawString(100, 650, "Languages: Python, Go, Java, TypeScript")
    c.drawString(100, 635, "Frameworks: FastAPI, Django, Spring Boot")
    c.drawString(100, 620, "Databases: PostgreSQL, Redis, Cassandra")
    c.drawString(100, 605, "Cloud & DevOps: AWS, Docker, Kubernetes, Terraform")
    c.drawString(100, 590, "Tools: Git, Kafka, Prometheus")
    c.drawString(100, 570, "Work Experience:")
    c.drawString(100, 555, "Senior Software Engineer - Stripe (2021 - Present)")
    c.drawString(100, 540, "Engineered fault-tolerant distributed payment settlement pipelines using Kafka.")
    c.drawString(100, 520, "Software Engineer - Twilio (2018 - 2021)")
    c.drawString(100, 505, "Built scalable messaging microservices in Python and AWS ECS.")
    c.drawString(100, 485, "Education:")
    c.drawString(100, 470, "B.S. in Computer Science - University of California, Berkeley (2014 - 2018)")
    c.drawString(100, 450, "Projects:")
    c.drawString(100, 435, "Raft Consensus Cluster: Implemented distributed consensus protocol in Go.")
    c.save()
    buf.seek(0)
    return buf.read()


def create_sample_docx() -> bytes:
    doc = Document()
    doc.add_heading("Alex Rivera", 0)
    doc.add_paragraph("alex.rivera@techcorp.io | +1-202-555-0144 | Austin, TX")
    doc.add_paragraph("Senior Cloud Architect with 10 years experience.")
    doc.add_heading("Skills", level=1)
    doc.add_paragraph("Languages: Go, Python, Rust")
    doc.add_paragraph("Frameworks: Gin, FastAPI")
    doc.add_paragraph("Databases: PostgreSQL, DynamoDB")
    doc.add_paragraph("Cloud: GCP, Kubernetes, Terraform")
    doc.add_heading("Experience", level=1)
    doc.add_paragraph("Principal Architect at CloudScale (2020 - Present)")
    doc.add_heading("Education", level=1)
    doc.add_paragraph("M.S. Software Engineering - UT Austin")
    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf.read()


def test_01_pdf_text_extraction():
    pdf_bytes = create_sample_pdf()
    text = extract_text_from_pdf_bytes(pdf_bytes)
    assert "Jane Doe" in text
    assert "jane.doe@example.com" in text
    assert "Senior Backend Engineer" in text
    assert "FastAPI" in text


def test_02_docx_text_extraction():
    docx_bytes = create_sample_docx()
    text = extract_text_from_docx_bytes(docx_bytes)
    assert "Alex Rivera" in text
    assert "alex.rivera@techcorp.io" in text
    assert "CloudScale" in text


def test_03_txt_text_extraction():
    txt_content = "Sam Smith\nsam@example.com\nPython, Docker, SQL\n5 years experience"
    text = extract_text_from_txt_bytes(txt_content.encode("utf-8"))
    assert text == txt_content


def test_04_candidate_claims_parsing():
    pdf_bytes = create_sample_pdf()
    claims = parse_resume_from_bytes(pdf_bytes, "jane_doe_resume.pdf")
    assert isinstance(claims, CandidateClaims)
    assert claims.name == "Jane Doe"
    assert claims.email == "jane.doe@example.com"
    assert claims.phone == "+1-555-0199"
    assert "San Francisco" in claims.location
    assert "Backend" in claims.current_role
    assert "Python" in claims.claimed_languages
    assert "FastAPI" in claims.claimed_frameworks
    assert "PostgreSQL" in claims.claimed_databases
    assert "AWS" in claims.claimed_cloud_devops
    assert len(claims.education) >= 1
    assert len(claims.projects) >= 1


def test_05_zero_hallucination_for_missing_fields():
    # Valid resume structure but missing email, phone, location, and certifications
    minimal_resume = (
        "John NoContact\n"
        "Professional Summary:\n"
        "Software engineer with background in distributed architectures.\n"
        "Technical Skills:\n"
        "Languages: Python, Go, Rust\n"
        "Work Experience:\n"
        "Software Developer - OpenSource Org (2020 - 2023)\n"
        "Contributed to core distributed consensus modules.\n"
    )
    claims = parse_resume_from_bytes(minimal_resume.encode("utf-8"), "minimal.txt")
    assert claims.name == "John NoContact"
    assert claims.email is None
    assert claims.phone is None
    assert claims.location is None
    assert claims.certifications == []
    assert "Python" in claims.claimed_languages
    assert "Go" in claims.claimed_languages


@pytest.mark.asyncio
async def test_06_upload_endpoint_happy_path():
    transport = ASGITransport(app=app)
    async with AsyncClient(transport=transport, base_url="http://test") as ac:
        job_payload = {
            "title": "Staff Backend Platform Engineer",
            "department": "Platform Core",
            "raw_jd_text": "Building next-generation distributed database engines with Go, Python, and PostgreSQL.",
            "min_years_experience": 6.0
        }
        create_job_res = await ac.post("/api/v1/jobs", json=job_payload)
        assert create_job_res.status_code == 200
        job_data = create_job_res.json()
        job_id = job_data["id"]

        pdf_bytes = create_sample_pdf()
        files = {
            "file": ("jane_doe_resume.pdf", pdf_bytes, "application/pdf")
        }
        response = await ac.post(f"/api/v1/jobs/{job_id}/candidates/upload", files=files)
        assert response.status_code == 201
        cand_data = response.json()

        # Verify candidate response structure
        assert cand_data["candidate_id"] is not None
        import uuid as _uuid
        assert _uuid.UUID(cand_data["candidate_id"])
        assert cand_data["associated_job"]["id"] == job_id
        assert cand_data["associated_job"]["title"] == "Staff Backend Platform Engineer"
        assert cand_data["candidate"]["name"] == "Jane Doe"
        assert cand_data["candidate"]["email"] == "jane.doe@example.com"
        assert cand_data["candidate"]["phone"] == "+1-555-0199"

        # Verify claims categorization
        tech_claims = cand_data["parsed_resume_claims"]["technical_skills"]
        assert "Python" in tech_claims["languages"]
        assert "FastAPI" in tech_claims["frameworks"]
        assert "PostgreSQL" in tech_claims["databases"]

        # Verify STRICT BOUNDARIES: zero qualification scoring or recommendations
        assert cand_data["scores"]["qualification_score"] is None
        assert cand_data["scores"]["github_audit_score"] is None
        assert cand_data["recommendation"] is None
        assert "claims" in cand_data["verification_status"].lower() or "unverified" in cand_data["verification_status"].lower()


@pytest.mark.asyncio
async def test_07_job_candidates_list_and_detail_retrieval():
    transport = ASGITransport(app=app)
    async with AsyncClient(transport=transport, base_url="http://test") as ac:
        job_res = await ac.post("/api/v1/jobs", json={
            "title": "Cloud Infrastructure Architect",
            "department": "Infrastructure",
            "raw_jd_text": "Architecting modern multi-region cloud infrastructures with AWS, Terraform, and Kubernetes."
        })
        assert job_res.status_code == 200
        job_id = job_res.json()["id"]

        docx_bytes = create_sample_docx()
        files = {"file": ("alex_rivera_cv.docx", docx_bytes, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")}
        upload_res = await ac.post(f"/api/v1/jobs/{job_id}/candidates/upload", files=files)
        assert upload_res.status_code == 201
        cand_id = upload_res.json()["candidate_id"]

        list_res = await ac.get(f"/api/v1/jobs/{job_id}/candidates")
        assert list_res.status_code == 200
        candidates = list_res.json()["candidates"]
        assert any(c["candidate_id"] == cand_id for c in candidates)

        detail_res = await ac.get(f"/api/v1/candidates/{cand_id}")
        assert detail_res.status_code == 200
        detail = detail_res.json()
        assert detail["candidate_id"] == cand_id
        assert detail["associated_job"]["id"] == job_id
        assert detail["candidate"]["name"] == "Alex Rivera"
        assert "GCP" in detail["parsed_resume_claims"]["technical_skills"]["cloud_and_devops"]


@pytest.mark.asyncio
async def test_08_upload_rejection_on_invalid_file_format():
    transport = ASGITransport(app=app)
    async with AsyncClient(transport=transport, base_url="http://test") as ac:
        job_res = await ac.post("/api/v1/jobs", json={"title": "Test Engineer", "raw_jd_text": "Test engineering role."})
        assert job_res.status_code == 200
        job_id = job_res.json()["id"]

        files = {"file": ("resume.png", b"\x89PNG\r\n\x1a\nfakeimagecontent", "image/png")}
        res = await ac.post(f"/api/v1/jobs/{job_id}/candidates/upload", files=files)
        assert res.status_code == 400
        assert "unsupported" in res.json()["detail"].lower()


@pytest.mark.asyncio
async def test_09_upload_rejection_on_empty_file():
    transport = ASGITransport(app=app)
    async with AsyncClient(transport=transport, base_url="http://test") as ac:
        job_res = await ac.post("/api/v1/jobs", json={"title": "Test Engineer 2", "raw_jd_text": "Test role description."})
        assert job_res.status_code == 200
        job_id = job_res.json()["id"]

        files = {"file": ("empty.pdf", b"", "application/pdf")}
        res = await ac.post(f"/api/v1/jobs/{job_id}/candidates/upload", files=files)
        assert res.status_code == 400
        assert "empty" in res.json()["detail"].lower()


@pytest.mark.asyncio
async def test_10_upload_rejection_on_corrupt_file():
    transport = ASGITransport(app=app)
    async with AsyncClient(transport=transport, base_url="http://test") as ac:
        job_res = await ac.post("/api/v1/jobs", json={"title": "Test Engineer 3", "raw_jd_text": "Test role description."})
        assert job_res.status_code == 200
        job_id = job_res.json()["id"]

        files = {"file": ("corrupt.pdf", b"not really a pdf file at all", "application/pdf")}
        res = await ac.post(f"/api/v1/jobs/{job_id}/candidates/upload", files=files)
        assert res.status_code == 400
        assert "corrupt" in res.json()["detail"].lower() or "invalid" in res.json()["detail"].lower() or "failed" in res.json()["detail"].lower()


@pytest.mark.asyncio
async def test_11_nonexistent_job_returns_404():
    transport = ASGITransport(app=app)
    async with AsyncClient(transport=transport, base_url="http://test") as ac:
        files = {"file": ("sample.txt", b"Bob Tester\nSkill: Python\nExperience: Software Engineer 5 years", "text/plain")}
        res = await ac.post("/api/v1/jobs/00000000-0000-0000-0000-000000000000/candidates/upload", files=files)
        assert res.status_code == 404
        assert "not found" in res.json()["detail"].lower()
